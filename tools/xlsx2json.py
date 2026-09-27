"""Exporta las hojas de los Excel originales a JSON (valores calculados) para probar los lectores.
Las fechas se exportan como {"$d": "AAAA-MM-DD"} y las horas como {"$t": "HH:MM"}, imitando
lo que Apps Script entrega con getValues()."""
import datetime, json, pathlib, sys
import openpyxl

def cell(v):
    if isinstance(v, datetime.datetime):
        if v.year < 1901:
            return {"$t": v.strftime("%H:%M")}
        return {"$d": v.strftime("%Y-%m-%d")}
    if isinstance(v, datetime.time):
        return {"$t": v.strftime("%H:%M")}
    if isinstance(v, datetime.date):
        return {"$d": v.isoformat()}
    return "" if v is None else v

root = pathlib.Path(__file__).resolve().parent.parent
out = root / "tools" / "fixtures"
out.mkdir(parents=True, exist_ok=True)
for f in sorted((root / "originales").glob("*.xlsx")):
    wb = openpyxl.load_workbook(f, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        values = [[cell(c.value) for c in row] for row in ws.iter_rows(max_row=min(ws.max_row, 400))]
        while values and all(v == "" for v in values[-1]):
            values.pop()
        sheets.append({"name": ws.title, "hidden": ws.sheet_state != "visible", "values": values})
    (out / (f.stem + ".json")).write_text(json.dumps({"file": f.name, "sheets": sheets}, ensure_ascii=False))
    print("ok", f.name, len(sheets), "hojas")
